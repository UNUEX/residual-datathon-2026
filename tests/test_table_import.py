from io import BytesIO
import base64
import unittest
from unittest.mock import patch
import zipfile
import zlib

from src import table_import
from src.table_import import ImportError, inspect_file, read_table


class TextImportTests(unittest.TestCase):
    def test_csv_quotes_decimal_comma_multiline_and_identifier_preserved(self):
        content = 'Код;Площадь;Описание\r\n001;"12,50";"Город; район"\r\n002;20;"две\nстроки"\r\n'.encode("utf-8-sig")
        preview = inspect_file("sample.CSV", content)
        self.assertEqual(preview["format"], "csv")
        self.assertEqual(preview["tables"][0]["row_count"], 2)
        result = read_table("sample.csv", content, "csv:0")
        self.assertEqual(result["columns"], ["Код", "Площадь", "Описание"])
        self.assertEqual(result["rows"][0], ["001", "12,50", "Город; район"])
        self.assertEqual(result["rows"][1][2], "две\nстроки")

    def test_tsv_utf16_and_cp1251(self):
        text = "Код\tГород\n0001\tАстана\n"
        result = read_table("sample.tsv", text.encode("utf-16"), "csv:0")
        self.assertEqual(result["rows"], [["0001", "Астана"]])
        cp = read_table("sample.csv", text.replace("\t", ",").encode("cp1251"), "csv:0")
        self.assertEqual(cp["rows"], [["0001", "Астана"]])
        self.assertIn("Windows-1251", cp["warnings"][0])

    def test_header_selection_blank_rows_and_explicit_options(self):
        content = b"Survey 2024\n\nID;Value\n001;10\n\n002;20\n"
        result = read_table("sample.csv", content, "csv:0", header_row=3, options={"delimiter": ";", "encoding": "utf-8"})
        self.assertEqual(result["columns"], ["ID", "Value"])
        self.assertEqual(len(result["rows"]), 2)
        self.assertTrue(result["warnings"])

    def test_malformed_headers_rows_and_quotes_rejected(self):
        examples = [b"ID,\n001,2\n", b"id,ID\n001,2\n", b"ID,Value\n001,2,3\n",
                    b'ID;Value\n001;"6"0\n', b'ID,Value\n001,"unclosed\n']
        for content in examples:
            with self.subTest(content=content), self.assertRaises(ImportError):
                read_table("sample.csv", content, "csv:0")

    def test_single_column_and_literal_formula_are_plain_text(self):
        result = read_table("sample.csv", b"ID\n0001\n=1+1\n", "csv:0")
        self.assertEqual(result["rows"], [["0001"], ["=1+1"]])

    def test_invalid_inputs_and_limits(self):
        for name, content in [("script.py", b"print('hello')"), ("empty.csv", b""), ("sample.csv", b"A\x00B\n1\n")]:
            with self.subTest(name=name), self.assertRaises(ImportError):
                inspect_file(name, content)
        for header in [0, 4, True, "1"]:
            with self.subTest(header=header), self.assertRaises(ImportError):
                read_table("sample.csv", b"A\n1\n", "csv:0", header_row=header)
        with self.assertRaises(ImportError):
            read_table("sample.csv", b"A\n1\n", "missing")
        with self.assertRaises(ImportError):
            read_table("sample.csv", b"A\n1\n", "csv:0", options={"delimiter": "|"})
        with patch.object(table_import, "MAX_FILE_BYTES", 4), self.assertRaises(ImportError):
            inspect_file("sample.csv", b"ID\n01\n")
        with patch.object(table_import, "MAX_ROWS", 1), self.assertRaises(ImportError):
            inspect_file("sample.csv", b"ID\n01\n02\n")
        with patch.object(table_import, "MAX_COLUMNS", 1), self.assertRaises(ImportError):
            inspect_file("sample.csv", b"ID,Value\n01,2\n")
        with self.assertRaises(ImportError):
            inspect_file("sample.csv", ("A\n" + "x" * 2001).encode())


class HtmlImportTests(unittest.TestCase):
    def test_actual_tables_only_no_code_or_external_resources(self):
        content = b'''<html><script>throw new Error("DO NOT RUN");</script>
        <img src="https://invalid.example/file"><table><caption>First</caption>
        <tr><th>ID</th><th>Value</th></tr><tr><td>001</td><td><b>2</b></td></tr></table>
        <table><tr><th>Name</th><th>Note</th></tr><tr><td>A&amp;B</td>
        <td>yes<script>fetch("https://invalid.example")</script><br>ok</td></tr></table></html>'''
        with patch("urllib.request.urlopen", side_effect=AssertionError("network must not be called")):
            preview = inspect_file("tables.html", content)
            first = read_table("tables.html", content, "html:0")
            second = read_table("tables.html", content, "html:1")
        self.assertEqual(len(preview["tables"]), 2)
        self.assertEqual(preview["tables"][0]["label"], "First")
        self.assertEqual(first["rows"], [["001", "2"]])
        self.assertEqual(second["rows"], [["A&B", "yes ok"]])
        self.assertTrue(second["warnings"])

    def test_reject_nested_merged_missing_and_unclosed_tables(self):
        examples = ["<table><tr><td><table></table></td></tr></table>",
                    '<table><tr><th colspan="2">X</th></tr></table>',
                    '<table><tr><th rowspan="0">X</th></tr></table>',
                    "<table><tr><th>X</th></tr>", "<p>No table</p>",
                    "<table><tr><td>X</tr></table>"]
        for text in examples:
            with self.subTest(text=text), self.assertRaises(ImportError):
                inspect_file("sample.html", text.encode())

    def test_reject_nonrectangular_html(self):
        text = "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td></tr></table>"
        with self.assertRaises(ImportError):
            read_table("sample.htm", text.encode(), "html:0")


class ExcelImportTests(unittest.TestCase):
    # A small BIFF8 workbook generated with xlwt 1.3.0: value, formula and
    # merged sheets. Keeping the fixture inline avoids an XLS writer dependency.
    XLS_FIXTURE = "c-rk)O=uHA6#iy6|D<V>ra~)*P_Sy-gFyr@ZB%NDhnCudiU^vTAks9Er9J3Dt2HMP1Wy((g?jK(_2Liypf?XfMGt}?B#I|bqJoDu>zlU|O=E=yG!ofwGjHC!c{B4e@6FHdi^Wi3?sj_-Rp^6`lIex3LMAAA%Y`7Q-ZUj!&OMY&8w~}ohaBm0;z84MxAdX(cM%$HIi3;)-cdS?lNior)3G{YH`|aD(uQ8)qS(_YBCin*pp7jHGd;$%h3QS!e#mquOD-|(CEbEJgvG>g*A5lQ2el5a3lJ5!%WdAUB#u@*laYM(vaocW*kznZW*V#_>hNP8<kZ1*Hf_kV7uVqcwh{`6kgMnM)r<jO%4_;b%$C<zE<o~cE9LIr$UWq~YWWq796Am2sE}cthacM=vc`e_<YY3BW_)(avKs?AV`3^pq5x*qN*_ua$J41Lbn~t*^(g-L_7k%|nVCvkp%T|+7N>n-g4KQU5a3NCxMI_2;)UcVnxau9AL8W{B)gYxe0n)FHj+5Pl4(9Dtm9;(v>ztU$rjW&z%@EksoTZ0o9TJ(bv^U#Oe178<((s~iW!(;oeNysN#kr|yh?1U;5QSW|G4j7H2*%ah4_WR;_1l3>%<&7Xd+V7UOqGE5j|q|s>J<7ITvc%dCRozmo3Pn1(oflIy0NG3?H*0rdCPLj!^ny6@9__vd1%;kL?NJVvD7(Rp~!cHifz85JtKGJ}rzpQcPRnPt?PEn(aN-*LiL?DDPP@Z7M}KqC)<3b`xwj{W~|wxp_<$3Z?%6K4smHWwfe)Y}l}2!-kE<g-~|@onIQPw*VJkP^|Ad%als?XOjOGqsY-SNw1ZAkfq$fshZ+v9h{=vmsn$)^7lfC12npUag4DY<2A=G-4sRT|5<ax;GbJUwT@q^#m>Jx{&yl_H;M"

    @staticmethod
    def workbook_bytes(configure=None):
        import openpyxl
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Данные"
        sheet.append(["ID", "Area"])
        sheet.append(["0001", 12.5])
        if configure:
            configure(book)
        stream = BytesIO()
        book.save(stream)
        book.close()
        return stream.getvalue()

    def test_xlsx_sheets_strings_formatted_ids_and_header_choice(self):
        def configure(book):
            first = book.active
            first.append([2, 20])
            first["A3"].number_format = "0000"
            second = book.create_sheet("Другой набор")
            second.append(["Название обследования"])
            second.append(["Region", "Count"])
            second.append(["Астана", 42])
        content = self.workbook_bytes(configure)
        preview = inspect_file("book.xlsx", content)
        self.assertEqual([t["label"] for t in preview["tables"]], ["Данные", "Другой набор"])
        result = read_table("book.xlsx", content, "sheet:0")
        self.assertEqual(result["rows"], [["0001", "12.5"], ["0002", "20"]])
        second = read_table("book.xlsx", content, "sheet:1", header_row=2)
        self.assertEqual(second["rows"], [["Астана", "42"]])

    def test_selected_formula_sheet_blocked_other_sheet_still_readable(self):
        def configure(book):
            sheet = book.create_sheet("Формулы")
            sheet.append(["ID", "Value"])
            sheet.append(["001", "=1+1"])
        content = self.workbook_bytes(configure)
        preview = inspect_file("book.xlsx", content)
        self.assertTrue(preview["tables"][1]["warnings"])
        self.assertEqual(read_table("book.xlsx", content, "sheet:0")["rows"][0][0], "0001")
        with self.assertRaisesRegex(ImportError, "формулы"):
            read_table("book.xlsx", content, "sheet:1")

    def test_merged_and_error_cells_blocked(self):
        for configure in [lambda book: book.active.merge_cells("A3:B3"),
                          lambda book: setattr(book.active["B2"], "value", "#DIV/0!")]:
            with self.subTest(configure=configure), self.assertRaises(ImportError):
                read_table("book.xlsx", self.workbook_bytes(configure), "sheet:0")

    def test_zip_expansion_dtd_and_invalid_file_rejected(self):
        with patch.object(table_import, "MAX_EXPANDED_BYTES", 10), self.assertRaises(ImportError):
            inspect_file("book.xlsx", self.workbook_bytes())
        stream = BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("xl/worksheets/sheet1.xml", '<!DOCTYPE a [<!ENTITY a "b">]><a/>')
        with self.assertRaisesRegex(ImportError, "DTD"):
            inspect_file("book.xlsx", stream.getvalue())
        with self.assertRaises(ImportError):
            inspect_file("book.xlsx", b"not a workbook")

    def test_xls_values_preserved_and_formula_or_merged_sheet_rejected(self):
        content = zlib.decompress(base64.b85decode(self.XLS_FIXTURE))
        preview = inspect_file("book.xls", content)
        self.assertEqual([t["label"] for t in preview["tables"]], ["Data", "Formula", "Merged"])
        self.assertEqual(read_table("book.xls", content, "sheet:0")["rows"], [["0001", "12.5"], ["0002", "20"]])
        with self.assertRaisesRegex(ImportError, "формулы"):
            read_table("book.xls", content, "sheet:1")
        with self.assertRaisesRegex(ImportError, "объединённые"):
            read_table("book.xls", content, "sheet:2")

    def test_xlsx_wrong_declared_dimensions_cannot_silently_drop_data(self):
        original = BytesIO(self.workbook_bytes())
        changed = BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(changed, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'<dimension ref="A1:B2"/>', b'<dimension ref="A1:A1"/>')
                target.writestr(info, data)
        result = read_table("book.xlsx", changed.getvalue(), "sheet:0")
        self.assertEqual(result["columns"], ["ID", "Area"])
        self.assertEqual(result["rows"], [["0001", "12.5"]])


if __name__ == "__main__":
    unittest.main()
